# Фикстуры 05a открываются стандартными разборщиками (решение владельца по XLSX, D-024).
# Нужны openpyxl и python-docx: python3 artifacts/stage-05a/fixture-open-check.py <каталог выгрузки>.
import csv, importlib.metadata as md, sys, warnings
import docx, openpyxl

warnings.simplefilter('ignore')
d = sys.argv[1]
print(f"# openpyxl {openpyxl.__version__}, python-docx {md.version('python-docx')}, csv (стандартная библиотека Python)")
for name in ('smeta-stromynka.xlsx', 'edge.xlsx'):
    for data_only in (False, True):
        wb = openpyxl.load_workbook(f'{d}/{name}', data_only=data_only)
        print(f'{name} (data_only={data_only}): листы {wb.sheetnames}')
        for ws in wb.worksheets:
            print(f'  {ws.title} [{ws.sheet_state}]: {ws.dimensions}, объединено {[str(r) for r in ws.merged_cells.ranges]}')
            for row in ws.iter_rows():
                vals = [(c.coordinate, c.value) for c in row if c.value is not None]
                if vals:
                    print('   ', vals)
doc = docx.Document(f'{d}/contract.docx')
print(f'contract.docx: абзацев {len(doc.paragraphs)}, таблиц {len(doc.tables)}, разделов {len(doc.sections)}')
for p in doc.paragraphs:
    print('   ', p.text)
for t in doc.tables:
    for r in t.rows:
        print('    |', ' | '.join(c.text for c in r.cells))
with open(f'{d}/cp1251.csv', encoding='cp1251', newline='') as f:
    rows = list(csv.reader(f, delimiter=';'))
print(f'cp1251.csv (cp1251, «;»): записей {len(rows)}')
for r in rows:
    print('   ', r)
